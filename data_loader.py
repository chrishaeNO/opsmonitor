from __future__ import annotations

from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple
from datetime import date, datetime
from pathlib import Path
import csv
import re
import tempfile
import unicodedata
import urllib.request
from urllib.parse import urlparse
from urllib.error import URLError

try:
    import openpyxl
    EXCEL_SUPPORT = True
except ImportError:
    EXCEL_SUPPORT = False

from models import PersonnelRecord

MAX_LOCAL_FILE_BYTES = 50 * 1024 * 1024
MAX_DOWNLOAD_BYTES = 50 * 1024 * 1024
URL_TIMEOUT_SECONDS = 20
ALLOWED_SUFFIXES = {".xlsx", ".csv"}


FIELD_SYNONYMS = {
    "name": ["navn", "name", "full name", "fullname"],
    "radio": ["sambandsnummer", "samband", "radio", "radio nr", "radio nummer"],
    "team": ["team", "gruppe", "unit"],
    "area": ["område", "omrade", "area", "sone", "zone"],
    "planned_on": ["planlagtpåtropp", "planlagtpatropp", "planlagt påtropp", "planned_on", "planned on", "start"],
    "planned_off": ["planlagtavtropp", "planlagt avtropp", "planned_off", "planned off", "end"],
    "actual_on": ["skannetpåtropp", "skannetpatropp", "skannet påtropp", "actual_on", "actual on", "innskann"],
    "actual_off": ["avtropp", "skannet avtropp", "actual_off", "actual off", "utskann"],
}


def is_url(source: str) -> bool:
    parsed = urlparse(source)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def sharepoint_download_url(url: str) -> str:
    if "download=" in url:
        return url
    separator = "&" if "?" in url else "?"
    return f"{url}{separator}download=1"

@dataclass(frozen=True)
class ResolvedSource:
    path: str
    is_temp: bool = False


def _normalize_header(value: str) -> str:
    s = (value or "").strip().replace("\ufeff", "").replace("\u00a0", " ")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.lower()
    s = re.sub(r"[\s\-_./]+", " ", s)
    s = re.sub(r"[^\w ]+", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _normalized_keys(value: str) -> Tuple[str, str]:
    n = _normalize_header(value)
    return n, n.replace(" ", "")


def _safe_suffix_from_url(url: str) -> str:
    suffix = Path(urlparse(url).path).suffix.lower()
    return suffix if suffix in ALLOWED_SUFFIXES else ""


def _sniff_kind(path: Path) -> str:
    try:
        with path.open("rb") as f:
            head = f.read(8)
    except Exception:
        return "unknown"
    if head.startswith(b"PK"):
        return "xlsx"
    try:
        head.decode("utf-8")
        return "text"
    except Exception:
        return "binary"


def _validate_local_file(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")
    if not path.is_file():
        raise ValueError(f"Path is not a file: {path}")
    if path.suffix.lower() not in ALLOWED_SUFFIXES:
        raise ValueError(f"Unsupported file format: {path.suffix.lower()}")
    try:
        size = path.stat().st_size
    except Exception:
        size = None
    if size is not None and size > MAX_LOCAL_FILE_BYTES:
        raise ValueError(f"File too large ({size} bytes). Max allowed is {MAX_LOCAL_FILE_BYTES} bytes.")


def _download_to_temp(url: str) -> ResolvedSource:
    url = sharepoint_download_url(url)
    suffix = _safe_suffix_from_url(url) or ".xlsx"
    fd, temp_path = tempfile.mkstemp(prefix="ops_sharepoint_", suffix=suffix)
    try:
        import os
        os.close(fd)
    except Exception:
        pass

    p = Path(temp_path)
    try:
        p.unlink(missing_ok=True)
        req = urllib.request.Request(url, headers={"User-Agent": "OPS/1.0"})
        with urllib.request.urlopen(req, timeout=URL_TIMEOUT_SECONDS) as resp:
            total = 0
            with p.open("wb") as out:
                while True:
                    chunk = resp.read(64 * 1024)
                    if not chunk:
                        break
                    total += len(chunk)
                    if total > MAX_DOWNLOAD_BYTES:
                        raise ValueError(f"Download too large (>{MAX_DOWNLOAD_BYTES} bytes).")
                    out.write(chunk)
    except URLError as exc:
        p.unlink(missing_ok=True)
        raise ValueError(f"Could not download URL: {exc}") from exc
    except Exception:
        p.unlink(missing_ok=True)
        raise

    kind = _sniff_kind(p)
    if suffix == ".xlsx" and kind != "xlsx":
        p.unlink(missing_ok=True)
        raise ValueError("Downloaded file is not a valid .xlsx (expected ZIP container).")
    if suffix == ".csv" and kind not in {"text", "unknown"}:
        p.unlink(missing_ok=True)
        raise ValueError("Downloaded file is not a valid text/CSV file.")
    return ResolvedSource(path=str(p), is_temp=True)


def resolve_source(source: str) -> ResolvedSource:
    if not source:
        return ResolvedSource(path=source, is_temp=False)
    if is_url(source):
        return _download_to_temp(source)
    path = Path(source)
    _validate_local_file(path)
    if path.suffix.lower() == ".xlsx" and _sniff_kind(path) != "xlsx":
        raise ValueError("File is not a valid .xlsx (expected ZIP container).")
    return ResolvedSource(path=str(path), is_temp=False)


def source_to_local_path(source: str) -> str:
    # Backwards compatible helper (older callers). Prefer `resolve_source`.
    return resolve_source(source).path


def parse_datetime(value) -> Optional[datetime]:
    if not value or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)) and EXCEL_SUPPORT:
        try:
            # Excel serial date/time (typical range for modern dates)
            if 20000 <= float(value) <= 90000:
                from openpyxl.utils.datetime import from_excel
                return from_excel(value)
        except Exception:
            pass
    
    formats = [
        "%Y-%m-%d %H:%M:%S",
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y %H:%M",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%Y-%m-%d %H:%M",
    ]
    
    for fmt in formats:
        try:
            return datetime.strptime(str(value).strip(), fmt)
        except ValueError:
            continue
    
    return None


def load_excel_data(file_path: str, sheet_index: int, column_mapping: dict) -> List[PersonnelRecord]:
    if not EXCEL_SUPPORT:
        raise ImportError("openpyxl is not installed. Install it with: pip install openpyxl")
    
    path = Path(file_path)
    _validate_local_file(path)
    
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.worksheets[sheet_index] if sheet_index < len(wb.worksheets) else wb.active
    
    headers = []
    for cell in sheet[1]:
        headers.append(str(cell.value).strip() if cell.value else "")

    effective_mapping = infer_column_mapping(headers, column_mapping)
    
    normalized_to_index: Dict[str, int] = {}
    for idx, h in enumerate(headers):
        if not h:
            continue
        k1, k2 = _normalized_keys(h)
        normalized_to_index[k1] = idx
        normalized_to_index[k2] = idx

    col_indices: Dict[str, int] = {}
    for field, col_name in effective_mapping.items():
        if not col_name:
            col_indices[field] = -1
            continue
        k1, k2 = _normalized_keys(col_name)
        idx = normalized_to_index.get(k1, normalized_to_index.get(k2, -1))
        col_indices[field] = idx
    
    records = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        
        name = str(row[col_indices["name"]]).strip() if col_indices["name"] >= 0 and row[col_indices["name"]] else ""
        if not name:
            continue
        
        planned_off_idx = col_indices.get("planned_off", -1)
        record = PersonnelRecord(
            name=name,
            radio=str(row[col_indices["radio"]]).strip() if col_indices["radio"] >= 0 and row[col_indices["radio"]] else "",
            team=str(row[col_indices["team"]]).strip() if col_indices["team"] >= 0 and row[col_indices["team"]] else "",
            area=str(row[col_indices["area"]]).strip() if col_indices["area"] >= 0 and row[col_indices["area"]] else "",
            planned_on=parse_datetime(row[col_indices["planned_on"]]) if col_indices["planned_on"] >= 0 else None,
            planned_off=parse_datetime(row[planned_off_idx]) if planned_off_idx >= 0 and row[planned_off_idx] else None,
            actual_on=parse_datetime(row[col_indices["actual_on"]]) if col_indices["actual_on"] >= 0 else None,
            actual_off=parse_datetime(row[col_indices["actual_off"]]) if col_indices["actual_off"] >= 0 else None,
        )
        records.append(record)
    
    return records


def load_csv_data(file_path: str, column_mapping: dict) -> List[PersonnelRecord]:
    path = Path(file_path)
    _validate_local_file(path)
    
    with open(file_path, "r", encoding="utf-8-sig", newline="") as f:
        # Prøv å autodetektere delimiter (Forms bruker ofte semikolon)
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,")
        except Exception:
            dialect = csv.excel

        reader = csv.DictReader(f, dialect=dialect)
        headers = reader.fieldnames or []
        # Bruk samme synonymlogikk som for Excel, kombinert med config-kolonner
        effective_mapping = infer_column_mapping(headers, column_mapping)

        records: List[PersonnelRecord] = []
        for row in reader:
            # Normaliser nøkler slik at vi ikke bryr oss om mellomrom/æøå-varianter
            normalized_row = { _normalize_header(k): v for k, v in row.items() if k }

            def get_value(field_key: str) -> str:
                col_name = effective_mapping.get(field_key, "")
                if not col_name:
                    return ""
                norm_key = _normalize_header(col_name)
                return (normalized_row.get(norm_key) or "").strip()

            name = get_value("name")
            if not name:
                continue

            record = PersonnelRecord(
                name=name,
                radio=get_value("radio"),
                team=get_value("team"),
                area=get_value("area"),
                planned_on=parse_datetime(get_value("planned_on")),
                planned_off=parse_datetime(get_value("planned_off")),
                actual_on=parse_datetime(get_value("actual_on")),
                actual_off=parse_datetime(get_value("actual_off")),
            )
            records.append(record)

    return records


def load_data_from_file(file_path: str, sheet_index: int, column_mapping: dict) -> List[PersonnelRecord]:
    resolved = resolve_source(file_path)
    try:
        path = Path(resolved.path)
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            return load_excel_data(resolved.path, sheet_index, column_mapping)
        if suffix == ".csv":
            return load_csv_data(resolved.path, column_mapping)
        raise ValueError(f"Unsupported file format: {suffix}")
    finally:
        if resolved.is_temp:
            try:
                Path(resolved.path).unlink(missing_ok=True)
            except Exception:
                pass


def get_excel_headers(file_path: str, sheet_index: int = 0) -> List[str]:
    if not EXCEL_SUPPORT:
        raise ImportError("openpyxl is not installed. Install it with: pip install openpyxl")

    resolved = resolve_source(file_path)
    try:
        wb = openpyxl.load_workbook(resolved.path, data_only=True, read_only=True)
        sheet = wb.worksheets[sheet_index] if sheet_index < len(wb.worksheets) else wb.active
        return [str(cell.value).strip() if cell.value else "" for cell in sheet[1]]
    finally:
        if resolved.is_temp:
            try:
                Path(resolved.path).unlink(missing_ok=True)
            except Exception:
                pass


def infer_column_mapping(headers: List[str], current: Optional[Dict[str, str]] = None) -> Dict[str, str]:
    mapping = dict(current or {})
    normalized: Dict[str, str] = {}
    for h in headers:
        if not h:
            continue
        k1, k2 = _normalized_keys(h)
        normalized[k1] = h
        normalized[k2] = h

    for field, candidates in FIELD_SYNONYMS.items():
        configured = mapping.get(field, "")
        if configured:
            k1, k2 = _normalized_keys(configured)
            if k1 in normalized or k2 in normalized:
                continue
            # Configured header isn't present in this sheet; try to infer instead.
            mapping[field] = ""
        for candidate in candidates:
            k1, k2 = _normalized_keys(candidate)
            if k1 in normalized:
                mapping[field] = normalized[k1]
                break
            if k2 in normalized:
                mapping[field] = normalized[k2]
                break

    for field in FIELD_SYNONYMS:
        if field not in mapping:
            mapping[field] = ""
    return mapping
